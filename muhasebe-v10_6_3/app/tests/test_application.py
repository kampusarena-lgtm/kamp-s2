import sys,shutil,tempfile,pathlib,unittest,threading,sqlite3,re
root=pathlib.Path(__file__).resolve().parents[1];target=pathlib.Path(tempfile.mkdtemp(prefix='kga-test-'));shutil.copytree(root,target,dirs_exist_ok=True)
sys.path.insert(0,str(target));import app as a
from datetime import date,timedelta
a.app.config['TESTING']=True;a.schedule_excel_sync=lambda:None
class ReviewTests(unittest.TestCase):
 def setUp(self):
  a.DB_FILE=target/('case-'+self._testMethodName+'.db');a.init_db();a.migrate_v105_personal_card_logic()
  self.client=a.app.test_client();self.client.get('/')
 def post(self,path,data=None,client=None):
  client=client or self.client
  with client.session_transaction() as s:token=s.get('csrf_token')
  response=client.post(path,data={'csrf_token':token,**(data or {})})
  return response
 def login(self):
  response=self.post('/admin-login',{'username':'Owner','password':'test-password','password_confirm':'test-password'});self.client.get('/admin');return response
 def staff(self):
  c=a.db_conn();x=c.execute("INSERT INTO personnel(name,monthly_salary,default_shift,active,created_at) VALUES('Personel',26000,'09-19',1,?)",(a.now_iso(),));pid=x.lastrowid;a.record_salary_rate(c,pid,26000)
  c.execute('INSERT INTO personnel_attendance(personnel_id,work_date,shift_code,scheduled_start,scheduled_end,actual_start,actual_end,worked_minutes,created_at,created_by) VALUES(?,?,?,?,?,?,?,?,?,?)',(pid,a.accounting_day_iso(),'09-19','09:00','19:00','09:00','19:00',600,a.now_iso(),'Personel'));c.commit();c.close();return pid
 def query(self,sql,args=()):
  c=a.db_conn();r=c.execute(sql,args).fetchall();c.close();return r
 def test_all_pages_and_post_csrf(self):
  self.assertEqual(self.login().status_code,302)
  for p in ['/', '/personel/mesai','/admin','/admin/cash','/admin/expenses','/admin/credit','/admin/partners','/admin/personnel','/admin/reports','/admin/day-close','/admin/audit','/admin/logs','/admin/settings']:
   with self.subTest(path=p):
    r=self.client.get(p);self.assertEqual(r.status_code,200)
    for form in re.findall(r'<form\b[^>]*method="post"[^>]*>(.*?)</form>',r.text,re.S):self.assertIn('name="csrf_token"',form)
  for p in (target/'templates').glob('*.html'):a.app.jinja_env.parse(p.read_text())
  self.assertEqual(self.client.post('/save-cash',data={}).status_code,400)
  self.assertEqual(self.client.get('/').headers['X-Frame-Options'],'DENY')
 def test_invalid_and_valid_amounts(self):
  for v in ['abc','-100','nan','inf','1.2345','1e3','1.23,45']:
   with self.subTest(amount=v):self.assertRaises(RuntimeError,a.money,v)
  self.assertEqual(a.money('1.234,56'),1234.56);self.assertEqual(a.money('1234.56'),1234.56)
  self.assertEqual(a.cash_calc(.1,.2,0,0,0,0,0),.3)
  self.login();self.post('/save-cash',{'tarih':a.accounting_day_iso(),'nakit':'abc','pos':'-100'})
  self.assertEqual(len(self.query('select * from cash_entries')),0)
  self.post('/save-cash',{'tarih':a.accounting_day_iso(),'nakit':'100.50','pos':'25'})
  self.assertEqual(self.query('select nakit from cash_entries')[0][0],100.5)
  self.post('/save-cash',{'tarih':a.accounting_day_iso(),'nakit':'50'})
  self.assertEqual(len(self.query('select * from cash_entries')),1)
 def test_old_key_rejected_and_role_loaded_from_db(self):
  self.login()
  from flask import Flask
  dummy=Flask('old');dummy.secret_key='kga-muhasebe-v10-change-me'
  cookie=dummy.session_interface.get_signing_serializer(dummy).dumps({'admin':True,'admin_id':1,'admin_owner':True})
  forged=a.app.test_client();forged.set_cookie('session',cookie)
  self.assertEqual(forged.get('/admin/settings').status_code,302)
  c=a.db_conn();c.execute('update admin_accounts set is_owner=0');c.commit();c.close()
  with self.client.session_transaction() as s:s['admin_owner']=True
  self.client.get('/admin/settings')
  with self.client.session_transaction() as s:self.assertFalse(s['admin_owner'])
  self.assertNotEqual(a.app.secret_key,dummy.secret_key)
 def test_first_setup_local_only(self):
  self.assertEqual(self.client.get('/admin-login',environ_overrides={'REMOTE_ADDR':'192.168.1.20'}).status_code,403)
  self.assertEqual(self.client.get('/admin-login').status_code,200)
 def test_concurrent_salary_and_meal(self):
  pid=self.staff();month=a.accounting_day().strftime('%Y-%m')
  for kind,expected in [('MAAŞ',1000),('YEMEK',200)]:
   barrier=threading.Barrier(2);results=[]
   def pay():
    barrier.wait()
    try:
     with a.app.test_request_context():results.append(a._create_personnel_payment(pid,month,kind,a.accounting_day_iso())[0])
    except RuntimeError as e:results.append(str(e))
   threads=[threading.Thread(target=pay) for _ in range(2)]
   for t in threads:t.start()
   for t in threads:t.join(10);self.assertFalse(t.is_alive())
   self.assertEqual(sum(isinstance(x,float) for x in results),1,results)
   self.assertEqual(self.query('select sum(amount) from personnel_payments where payment_kind=?',(kind,))[0][0],expected)
  self.assertEqual(self.query('select sum(amount) from expenses')[0][0],1200)
 def test_unique_cash_index(self):
  day=a.accounting_day_iso();c=a.db_conn()
  try:
   c.execute("INSERT INTO cash_entries(tx_date,created_at) VALUES(?,?)",(day,a.now_iso()));c.commit()
   with self.assertRaises(sqlite3.IntegrityError):c.execute("INSERT INTO cash_entries(tx_date,created_at) VALUES(?,?)",(day,a.now_iso()))
   c.rollback();c.execute('update cash_entries set active=0');c.commit();c.execute("INSERT INTO cash_entries(tx_date,created_at) VALUES(?,?)",(day,a.now_iso()));c.commit()
  finally:c.close()
 def test_past_salary_and_inactive_staff(self):
  pid=self.staff();month=a.accounting_day().strftime('%Y-%m');prev=(a.accounting_day().replace(day=1)-timedelta(days=1)).strftime('%Y-%m')
  c=a.db_conn();c.execute('insert into personnel_salary_history values(?,?,?)',(pid,prev,26000));c.execute("update personnel_attendance set work_date=?",(prev+'-15',));c.execute('update personnel set active=0,monthly_salary=52000');a.record_salary_rate(c,pid,52000);c.commit();c.close()
  r=a.personnel_month_summary(prev)[0];self.assertEqual(r['monthly_salary'],26000);self.assertEqual(r['salary_earned'],1000);self.assertEqual(r['active'],0)
 def test_expense_credit_and_personal_card_flow(self):
  self.login();today=a.accounting_day_iso()
  self.post('/save-expense',{'tarih':today,'kategori':'MARKET','tutar':'100','odeme_kaynagi':'İŞLETME'})
  self.assertEqual(a.profit_loss(today,today)['economic_expense'],100)
  self.post('/save-credit',{'tarih':today,'kisi':'Test','tur':'BORÇ','tutar':'50'})
  self.post('/save-credit',{'tarih':today,'kisi':'Test','tur':'ÖDEME','tutar':'20','yontem':'NAKİT'})
  self.assertEqual(self.query("select sum(case when movement_type='BORÇ' then amount else -amount end) from credit_entries")[0][0],30)
  self.post('/save-expense',{'tarih':today,'kategori':'MARKET','tutar':'200','odeme_kaynagi':'ŞAHSİ KART','kart_sahibi':'Murat'})
  self.assertEqual(a.pending_personal_card_total(today),200)
  pid=self.query('select id from partners where name=?',('Murat',))[0][0]
  self.post('/admin/partners/movement/add',{'date':today,'partner_id':pid,'type':'ALACAK ÖDEME','amount':'75','payment_method':'NAKİT','payment_target':'AUTO'})
  self.assertEqual(a.pending_personal_card_total(today),125)
  self.assertEqual(a.profit_loss(today,today)['economic_expense'],175)
  self.assertEqual(a.profit_loss(today,today)['general_cash_out'],175)
 def test_future_and_closed_attendance(self):
  self.staff();tomorrow=(a.accounting_day()+timedelta(days=1)).isoformat()
  self.assertRaises(RuntimeError,a.ensure_day_open,tomorrow)
  previous=(a.accounting_day()-timedelta(days=1)).isoformat()
  self.post('/personel/mesai/kaydet',{'personel':'Personel','tarih':previous,'vardiya':'09-19'})
  self.assertEqual(len(self.query('select * from personnel_attendance')),1)
 def test_connection_cleanup_after_integrity_failure(self):
  self.login();self.staff()
  self.post('/admin/settings/personnel/add',{'name':'Personel 2','monthly_salary':'100','default_shift':'09-19'})
  self.assertEqual(self.query('pragma integrity_check')[0][0],'ok')
 def test_payroll_rolls_back_when_expense_fails(self):
  pid=self.staff();c=a.db_conn();c.execute("CREATE TRIGGER fail_expense BEFORE INSERT ON expenses BEGIN SELECT RAISE(ABORT,'test rollback'); END");c.commit();c.close()
  with a.app.test_request_context():
   self.assertRaises(sqlite3.IntegrityError,a._create_personnel_payment,pid,a.accounting_day().strftime('%Y-%m'),'MAAŞ',a.accounting_day_iso())
  self.assertEqual(len(self.query('select * from personnel_payments')),0)
 def test_excel_exports_and_literal_notes(self):
  self.login();self.post('/save-expense',{'tarih':a.accounting_day_iso(),'kategori':'MARKET','tutar':'100','aciklama':'=1+1'})
  file=target/'export-test.xlsx';a.build_excel_mirror(file)
  from openpyxl import load_workbook
  wb=load_workbook(file);self.assertIn('MAAS_GECMISI',wb.sheetnames);self.assertEqual(wb['GIDERLER']['K2'].value,'=1+1');self.assertEqual(wb['GIDERLER']['K2'].data_type,'s');wb.close()
  response=self.post('/admin/export',{'start':a.accounting_day_iso(),'end':a.accounting_day_iso()})
  self.assertEqual(response.status_code,200);self.assertEqual(response.mimetype,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
 def test_migration_idempotent(self):
  self.staff();a.init_db();a.init_db();self.assertEqual(len(self.query('select * from personnel')),1)
if __name__=='__main__':unittest.main(verbosity=2)
